"""Word document loader (.docx) using python-docx."""

from pathlib import Path
from typing import Any, List, Optional

from src.modules.rag.domain.entities import Document
from src.modules.rag.domain.exceptions import DocumentParsingException
from src.modules.rag.domain.value_objects import DocumentType
from src.modules.rag.infrastructure.loaders.base_loader import BaseDocumentLoader


class DocxLoader(BaseDocumentLoader):
    """Extracts paragraphs and tables from Word (.docx) files."""

    supported_extensions = {"docx", "doc"}

    def load(
        self,
        file_path: str,
        tenant_id: Optional[str] = None,
        **kwargs: Any,
    ) -> Document:
        path = self.validate_file(file_path)

        try:
            # pyrefly: ignore [missing-import]
            import docx
        except ImportError as e:
            raise DocumentParsingException(
                "حزمة python-docx غير مثبتة. يرجى تثبيتها عبر: pip install python-docx"
            ) from e

        try:
            doc = docx.Document(str(path))
            content_parts: List[str] = []

            # Extract paragraphs
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    content_parts.append(text)

            # Extract tables as formatted text blocks
            tables_count = len(doc.tables)
            for table_idx, table in enumerate(doc.tables, 1):
                table_lines = [f"\n--- [جدول {table_idx}] ---"]
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    table_lines.append(" | ".join(row_cells))
                content_parts.append("\n".join(table_lines))

            metadata = {
                "paragraph_count": len(doc.paragraphs),
                "tables_count": tables_count,
            }

            full_content = "\n\n".join(content_parts)

            return self.create_document(
                path=path,
                content=full_content,
                doc_type=DocumentType.DOCX,
                metadata=metadata,
                tenant_id=tenant_id,
            )
        except Exception as err:
            raise DocumentParsingException(f"فشل في معالجة ملف الوورد (.docx): {err}") from err
