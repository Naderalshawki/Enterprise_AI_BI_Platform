"""Tests for document loaders."""

import tempfile
from pathlib import Path
# pyrefly: ignore [missing-import]
import pytest

from src.modules.rag.domain.entities import Document
from src.modules.rag.domain.exceptions import UnsupportedDocumentException
from src.modules.rag.domain.value_objects import DocumentType
from src.modules.rag.infrastructure.loaders.composite_loader import CompositeDocumentLoader
from src.modules.rag.infrastructure.loaders.text_loader import TextLoader


def test_text_loader_utf8():
    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".txt", encoding="utf-8") as f:
        f.write("هذا ملف نصي تجريبي باللغة العربية لاختبار مستخرج النصوص.")
        temp_path = f.name

    try:
        loader = TextLoader()
        assert loader.can_handle("txt")
        doc = loader.load(temp_path)

        assert isinstance(doc, Document)
        assert "ملف نصي تجريبي" in doc.text_content
        assert doc.doc_type == DocumentType.TXT
        assert doc.metadata["encoding"] in {"utf-8", "utf-8-sig"}
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_composite_loader_routing():
    composite = CompositeDocumentLoader()

    # Supported extensions
    assert composite.can_handle("pdf")
    assert composite.can_handle("docx")
    assert composite.can_handle("xlsx")
    assert composite.can_handle("csv")
    assert composite.can_handle("txt")
    assert composite.can_handle("md")

    # Unsupported extension
    assert not composite.can_handle("exe")

    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".unknown") as f:
        f.write("dummy")
        temp_path = f.name

    try:
        with pytest.raises(UnsupportedDocumentException):
            composite.load(temp_path)
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_csv_loader():
    from src.modules.rag.infrastructure.loaders.excel_loader import ExcelLoader

    csv_data = "الاسم,القسم,الراتب\nأحمد,المبيعات,8000\nسارة,التطوير,12000"
    with tempfile.NamedTemporaryFile("w", delete=False, suffix=".csv", encoding="utf-8") as f:
        f.write(csv_data)
        temp_path = f.name

    try:
        loader = ExcelLoader()
        assert loader.can_handle("csv")
        doc = loader.load(temp_path)

        assert isinstance(doc, Document)
        assert doc.doc_type == DocumentType.CSV
        assert "أحمد" in doc.text_content
        assert "سارة" in doc.text_content
        assert doc.metadata["sheets"][0]["rows"] == 2
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_docx_loader():
    import docx
    from src.modules.rag.infrastructure.loaders.docx_loader import DocxLoader

    with tempfile.NamedTemporaryFile("wb", delete=False, suffix=".docx") as f:
        temp_path = f.name

    try:
        d = docx.Document()
        d.add_heading("وثيقة تجريبية", level=1)
        d.add_paragraph("هذه فقرة تجريبية لاختبار مستخرج ملفات الوورد.")
        d.save(temp_path)

        loader = DocxLoader()
        assert loader.can_handle("docx")
        doc = loader.load(temp_path)

        assert isinstance(doc, Document)
        assert doc.doc_type == DocumentType.DOCX
        assert "وثيقة تجريبية" in doc.text_content
        assert "فقرة تجريبية" in doc.text_content
    finally:
        Path(temp_path).unlink(missing_ok=True)


def test_pdf_loader():
    import fitz
    from src.modules.rag.infrastructure.loaders.pdf_loader import PDFLoader

    with tempfile.NamedTemporaryFile("wb", delete=False, suffix=".pdf") as f:
        temp_path = f.name

    try:
        doc_pdf = fitz.open()
        page = doc_pdf.new_page()
        page.insert_text((50, 50), "PDF Test Content for RAG Document Extraction")
        doc_pdf.save(temp_path)
        doc_pdf.close()

        loader = PDFLoader()
        assert loader.can_handle("pdf")
        doc = loader.load(temp_path)

        assert isinstance(doc, Document)
        assert doc.doc_type == DocumentType.PDF
        assert "PDF Test Content" in doc.text_content
        assert doc.metadata["total_pages"] == 1
        assert len(doc.metadata["pages"]) == 1
        assert doc.metadata["pages"][0]["page_number"] == 1
    finally:
        Path(temp_path).unlink(missing_ok=True)


if __name__ == "__main__":
    pytest.main([__file__])
